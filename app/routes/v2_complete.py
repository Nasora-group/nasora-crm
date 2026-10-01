import io
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal
from flask import Blueprint, Response, jsonify, render_template, request, flash, redirect, url_for
from flask_login import current_user, login_required
from sqlalchemy import func, or_

from app.extensions import db
from app.models import (
    User, Prospection, Planning, AnimationSale, DIVISION_SUPPLIERS, SUPPLIERS,
    SalesObjective,
)
from app.models_clients import Client, ClientVisit
from app.models_stock import StockEntry
from app.models_v2_complete import (
    V2Opportunity, V2PlanningExecution, V2VisitGeo, V2Rupture,
    V2RestockRequest, V2Objective,
)
from app.permissions import is_admin, division_matches
from app.utils import decode_planning_slot, roles_required

v2_complete_bp = Blueprint("v2_complete", __name__, url_prefix="/v2")


def _divisions():
    return list(DIVISION_SUPPLIERS.keys())


def _division_arg():
    value = (request.args.get("division") or "all").strip().lower()
    return value if value in {"all", *_divisions()} else "all"


def _scope_user_query(query, division="all"):
    if not is_admin():
        query = query.filter(User.id == current_user.id)
    elif division in DIVISION_SUPPLIERS:
        query = query.filter(User.project == division)
    return query


def _period(year, month):
    start = date(year, month, 1)
    end = date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)
    previous = start - timedelta(days=1)
    prev_start = date(previous.year, previous.month, 1)
    return start, end, prev_start


def _revenue(division, start, end):
    total = Decimal("0")
    for slug in DIVISION_SUPPLIERS.get(division, []):
        model = SUPPLIERS[slug]["sale_model"]
        value = db.session.query(func.coalesce(func.sum(model.quantity * model.price), 0)).filter(
            model.project == division, model.date >= start, model.date < end
        ).scalar() or 0
        total += Decimal(str(value))
    return total


def _sync_planning_executions(start=None, end=None):
    start = start or date.today().replace(day=1)
    end = end or date.today() + timedelta(days=1)
    users = User.query.filter(User.role.in_(("commercial", "animateur"))).all()
    created = 0
    for planning in Planning.query.filter(Planning.date < end, Planning.date >= start - timedelta(days=7)).all():
        owner = next((u for u in users if u.id == planning.commercial_id), None)
        if not owner:
            continue
        for index, field in enumerate(("lundi", "mardi", "mercredi", "jeudi", "vendredi", "samedi", "dimanche")):
            execution_date = planning.date + timedelta(days=index)
            if execution_date < start or execution_date >= end:
                continue
            for structure_type, structure_name in decode_planning_slot(getattr(planning, field, None)):
                name = (structure_name or structure_type or "Structure non renseignée").strip()
                if not name:
                    continue
                existing = V2PlanningExecution.query.filter_by(
                    user_id=owner.id, execution_date=execution_date, structure_name=name
                ).first()
                if existing:
                    continue
                db.session.add(V2PlanningExecution(
                    planning_id=planning.id, user_id=owner.id, execution_date=execution_date,
                    weekday=field, structure_type=structure_type, structure_name=name, status="planned"
                ))
                created += 1
    if created:
        try:
            db.session.commit()
        except Exception:
            db.session.rollback()
    return created


def _automatic_evaluation(user, year, month):
    start, end, _ = _period(year, month)
    division = user.project
    target = Decimal(str(db.session.query(func.coalesce(func.sum(SalesObjective.target_amount), 0)).filter(
        SalesObjective.division == division, SalesObjective.year == year,
        SalesObjective.month == month
    ).scalar() or 0))
    ca = _revenue(division, start, end)
    if target:
        ca_user = Decimal("0")
        for slug in DIVISION_SUPPLIERS.get(division, []):
            model = SUPPLIERS[slug]["sale_model"]
            value = db.session.query(func.coalesce(func.sum(model.quantity * model.price), 0)).filter(
                model.commercial_id == user.id, model.date >= start, model.date < end,
                model.project == division
            ).scalar() or 0
            ca_user += Decimal(str(value))
    else:
        ca_user = Decimal("0")
    prospections = Prospection.query.filter_by(commercial_id=user.id).filter(
        Prospection.date >= start, Prospection.date < end
    ).count()
    visits = ClientVisit.query.filter_by(commercial_id=user.id, is_duplicate=False).filter(
        ClientVisit.date >= start, ClientVisit.date < end
    ).count()
    animations = AnimationSale.query.filter_by(animateur_id=user.id).filter(
        AnimationSale.animation_date >= start, AnimationSale.animation_date < end
    ).count()
    execution_total = V2PlanningExecution.query.filter_by(user_id=user.id).filter(
        V2PlanningExecution.execution_date >= start, V2PlanningExecution.execution_date < end
    ).count()
    execution_done = V2PlanningExecution.query.filter_by(user_id=user.id, status="realized").filter(
        V2PlanningExecution.execution_date >= start, V2PlanningExecution.execution_date < end
    ).count()
    ca_rate = float(ca_user / target * 100) if target else 0
    execution_rate = (execution_done / execution_total * 100) if execution_total else 0
    score = min(ca_rate, 100) * 0.45 + min(prospections / 50 * 100, 100) * 0.15 + min(visits / 50 * 100, 100) * 0.20 + execution_rate * 0.20
    level = "Excellent" if score >= 90 else "Bon" if score >= 75 else "Moyen" if score >= 60 else "À renforcer"
    return {
        "user": user, "target": target, "ca": ca_user, "ca_rate": ca_rate,
        "prospections": prospections, "visits": visits, "animations": animations,
        "planned": execution_total, "realized": execution_done,
        "execution_rate": execution_rate, "score": round(score, 1), "level": level
    }


@v2_complete_bp.route("/cockpit")
@login_required
@roles_required("admin")
def cockpit():
    today = date.today()
    year = request.args.get("year", today.year, type=int)
    month = request.args.get("month", today.month, type=int)
    division = _division_arg()
    start, end, prev_start = _period(year, month)
    prev_end = start
    _sync_planning_executions(start, end)

    divisions = _divisions() if division == "all" else [division]
    current_ca = sum((_revenue(d, start, end) for d in divisions), Decimal("0"))
    previous_ca = sum((_revenue(d, prev_start, prev_end) for d in divisions), Decimal("0"))
    targets = sum((Decimal(str(db.session.query(func.coalesce(func.sum(SalesObjective.target_amount), 0)).filter(
        SalesObjective.division == d, SalesObjective.year == year, SalesObjective.month == month
    ).scalar() or 0)) for d in divisions), Decimal("0"))

    users = User.query.filter(User.role.in_(("commercial", "animateur")), User.is_active_account.is_(True))
    if division != "all":
        users = users.filter(User.project == division)
    users = users.order_by(User.project, User.username).all()
    evaluations = [_automatic_evaluation(u, year, month) for u in users]

    visit_count = ClientVisit.query.join(User, ClientVisit.commercial_id == User.id).filter(
        ClientVisit.date >= start, ClientVisit.date < end, ClientVisit.is_duplicate.is_(False),
        User.project.in_(divisions)
    ).count()
    prospection_count = Prospection.query.join(User, Prospection.commercial_id == User.id).filter(
        Prospection.date >= start, Prospection.date < end, User.project.in_(divisions)
    ).count()
    executions = V2PlanningExecution.query.join(User, V2PlanningExecution.user_id == User.id).filter(
        V2PlanningExecution.execution_date >= start, V2PlanningExecution.execution_date < end,
        User.project.in_(divisions)
    ).all()
    animation_rows = AnimationSale.query.filter(
        AnimationSale.animation_date >= start, AnimationSale.animation_date < end,
        AnimationSale.project.in_(divisions)
    ).all()
    rupture_open = V2Rupture.query.filter(V2Rupture.division.in_(divisions), V2Rupture.status != "resolved").count()
    restock_open = V2RestockRequest.query.filter(V2RestockRequest.division.in_(divisions), V2RestockRequest.status != "resolved").count()
    opportunities = V2Opportunity.query.filter(V2Opportunity.division.in_(divisions)).count()
    followups = V2Opportunity.query.filter(
        V2Opportunity.division.in_(divisions), V2Opportunity.next_followup.isnot(None),
        V2Opportunity.next_followup <= date.today(), V2Opportunity.stage != "prescription_obtenue"
    ).count()

    return render_template(
        "v2/cockpit_complete.html",
        year=year, month=month, division=division, current_ca=current_ca,
        previous_ca=previous_ca, ca_evolution=((current_ca - previous_ca) / previous_ca * 100 if previous_ca else None),
        target=targets, target_rate=(current_ca / targets * 100 if targets else None),
        visits=visit_count, prospections=prospection_count,
        planned=len(executions), realized=sum(1 for e in executions if e.status == "realized"),
        execution_rate=(sum(1 for e in executions if e.status == "realized") / len(executions) * 100 if executions else 0),
        animation_count=len(animation_rows), animation_ca=sum((a.total_amount for a in animation_rows), Decimal("0")),
        rupture_open=rupture_open, restock_open=restock_open, opportunities=opportunities,
        followups=followups, evaluations=evaluations,
        periods={"current": start.isoformat(), "end": end.isoformat(), "previous": prev_start.isoformat()},
    )


@v2_complete_bp.route("/planning-execution", methods=["GET", "POST"])
@login_required
@roles_required("admin", "commercial", "animateur")
def planning_execution():
    if request.method == "POST":
        execution_id = request.form.get("execution_id", type=int)
        execution = V2PlanningExecution.query.get_or_404(execution_id)
        if not is_admin() and execution.user_id != current_user.id:
            return ("Forbidden", 403)
        status = (request.form.get("status") or "planned").strip()
        if status not in {"planned", "realized", "not_realized"}:
            status = "planned"
        execution.status = status
        execution.reason = (request.form.get("reason") or "").strip() or None
        execution.notes = (request.form.get("notes") or "").strip() or None
        execution.realized_at = datetime.utcnow() if status == "realized" else None
        db.session.commit()
        return redirect(url_for("v2_complete.planning_execution"))
    start = date.today() - timedelta(days=7)
    end = date.today() + timedelta(days=8)
    _sync_planning_executions(start, end)
    query = V2PlanningExecution.query.join(User, V2PlanningExecution.user_id == User.id).filter(
        V2PlanningExecution.execution_date >= start, V2PlanningExecution.execution_date < end
    )
    if not is_admin():
        query = query.filter(V2PlanningExecution.user_id == current_user.id)
    rows = query.order_by(V2PlanningExecution.execution_date.desc(), User.username).all()
    return render_template("v2/planning_execution.html", rows=rows, reasons=(
        "professionnel absent", "structure fermée", "déplacement", "rendez-vous annulé", "rupture de stock", "autre"
    ))


@v2_complete_bp.route("/opportunites", methods=["GET", "POST"])
@login_required
@roles_required("admin", "commercial", "animateur")
def opportunities():
    if request.method == "POST":
        opportunity_id = request.form.get("opportunity_id", type=int)
        client_id = request.form.get("client_id", type=int)
        client = Client.query.get_or_404(client_id)
        owner_id = client.owner_id or current_user.id
        if not is_admin() and owner_id != current_user.id:
            return ("Forbidden", 403)
        owner = User.query.get(owner_id)
        division = owner.project if owner else current_user.project
        stage = (request.form.get("stage") or "prospect").strip()
        stages = {"prospect", "produit_presente", "interet", "prescription_potentielle", "prescription_obtenue", "suivi"}
        if stage not in stages:
            flash("Étape du pipeline invalide.", "danger")
            return redirect(url_for("v2_complete.opportunities"))
        potential = request.form.get("potential_prescription", 0, type=int) or 0
        obtained = request.form.get("obtained_prescription", 0, type=int) or 0
        if potential < 0 or obtained < 0:
            flash("Les prescriptions doivent être positives.", "danger")
            return redirect(url_for("v2_complete.opportunities"))
        if stage == "prescription_obtenue" and obtained < 1:
            obtained = max(potential, 1)
        if opportunity_id:
            item = V2Opportunity.query.get_or_404(opportunity_id)
            if not is_admin() and item.owner_id != current_user.id:
                return ("Forbidden", 403)
            item.client_id = client.id
            item.division = division
        else:
            item = V2Opportunity(client_id=client.id, owner_id=owner_id, division=division)
            db.session.add(item)
        item.product_name = (request.form.get("product_name") or "").strip() or None
        item.stage = stage
        item.interest_level = (request.form.get("interest_level") or "").strip() or None
        item.potential_prescription = potential
        item.obtained_prescription = obtained
        item.next_followup = date.fromisoformat(request.form["next_followup"]) if request.form.get("next_followup") else None
        item.notes = (request.form.get("notes") or "").strip() or None
        db.session.commit()
        flash("Opportunité mise à jour." if opportunity_id else "Opportunité enregistrée.", "success")
        return redirect(url_for("v2_complete.opportunities"))
    query = V2Opportunity.query.join(Client, V2Opportunity.client_id == Client.id)
    if not is_admin():
        query = query.filter(V2Opportunity.owner_id == current_user.id)
    rows = query.order_by(V2Opportunity.next_followup.asc().nullslast(), V2Opportunity.updated_at.desc()).all()
    clients = Client.query.filter(Client.owner_id == current_user.id).order_by(Client.name).all() if not is_admin() else Client.query.order_by(Client.name).limit(500).all()
    return render_template("v2/opportunities.html", rows=rows, clients=clients,
        stages=("prospect", "produit_presente", "interet", "prescription_potentielle", "prescription_obtenue", "suivi"))
@v2_complete_bp.route("/visites-geolocalisees")
@login_required
@roles_required("admin", "commercial", "animateur")
def geo_visits():
    query = ClientVisit.query.join(User, ClientVisit.commercial_id == User.id).outerjoin(V2VisitGeo, V2VisitGeo.visit_id == ClientVisit.id).filter(
        ClientVisit.is_duplicate.is_(False), V2VisitGeo.id.isnot(None)
    )
    if not is_admin():
        query = query.filter(ClientVisit.commercial_id == current_user.id)
    rows = query.order_by(ClientVisit.date.desc()).limit(500).all()
    geos = {g.visit_id: g for g in V2VisitGeo.query.filter(V2VisitGeo.visit_id.in_([r.id for r in rows])).all()} if rows else {}
    return render_template("v2/geo_visits.html", rows=rows, geos=geos)


@v2_complete_bp.route("/visites/<int:visit_id>/geolocalisation", methods=["POST"])
@login_required
@roles_required("admin", "commercial", "animateur")
def save_visit_geo(visit_id):
    visit = ClientVisit.query.get_or_404(visit_id)
    if not is_admin() and visit.commercial_id != current_user.id:
        return jsonify({"error": "forbidden"}), 403
    payload = request.get_json(silent=True) or request.form
    try:
        lat = float(payload.get("latitude")); lng = float(payload.get("longitude"))
        accuracy = float(payload.get("accuracy")) if payload.get("accuracy") not in (None, "") else None
        duration = int(payload.get("duration_minutes")) if payload.get("duration_minutes") not in (None, "") else None
    except (TypeError, ValueError):
        return jsonify({"error": "coordonnées invalides"}), 400
    geo = V2VisitGeo.query.filter_by(visit_id=visit.id).first()
    if not geo:
        geo = V2VisitGeo(visit_id=visit.id, latitude=lat, longitude=lng, accuracy_m=accuracy, duration_minutes=duration)
        db.session.add(geo)
    else:
        geo.latitude, geo.longitude, geo.accuracy_m, geo.duration_minutes, geo.captured_at = lat, lng, accuracy, duration, datetime.utcnow()
    db.session.commit()
    return jsonify({"ok": True, "latitude": lat, "longitude": lng})


@v2_complete_bp.route("/ruptures", methods=["GET", "POST"])
@login_required
@roles_required("admin", "commercial", "animateur")
def ruptures():
    if request.method == "POST":
        division = (request.form.get("division") or current_user.project).lower()
        if not division_matches(current_user, division):
            return ("Forbidden", 403)
        row = V2Rupture(
            division=division, laboratory=request.form.get("laboratory") or None,
            product_name=request.form.get("product_name", "").strip(),
            wholesaler=request.form.get("wholesaler") or None,
            quantity=request.form.get("quantity", 0, type=int) or 0,
            reported_by_id=current_user.id, notes=request.form.get("notes") or None,
        )
        db.session.add(row); db.session.commit()
        return redirect(url_for("v2_complete.ruptures"))
    query = V2Rupture.query
    if not is_admin(): query = query.filter_by(division=current_user.project)
    rows = query.order_by(V2Rupture.reported_at.desc()).all()
    return render_template("v2/ruptures.html", rows=rows)


@v2_complete_bp.route("/ruptures/<int:rupture_id>/status", methods=["POST"])
@login_required
@roles_required("admin")
def rupture_status(rupture_id):
    row = V2Rupture.query.get_or_404(rupture_id)
    status = request.form.get("status") or "rupture"
    if status not in {"rupture", "procurement", "stock_available", "resolved"}:
        status = "rupture"
    row.status = status
    now = datetime.utcnow()
    if status == "procurement": row.procurement_at = now
    if status == "stock_available": row.stock_available_at = now
    if status == "resolved": row.resolved_at = now
    db.session.commit()
    return redirect(url_for("v2_complete.ruptures"))


@v2_complete_bp.route("/reassorts", methods=["GET", "POST"])
@login_required
@roles_required("admin", "commercial", "animateur")
def restocks():
    if request.method == "POST":
        division = (request.form.get("division") or current_user.project).lower()
        if not division_matches(current_user, division):
            return ("Forbidden", 403)
        row = V2RestockRequest(
            division=division, laboratory=request.form.get("laboratory") or None,
            product_name=request.form.get("product_name", "").strip(),
            wholesaler=request.form.get("wholesaler") or None,
            quantity=request.form.get("quantity", 1, type=int) or 1,
            requested_by_id=current_user.id, notes=request.form.get("notes") or None,
        )
        db.session.add(row); db.session.commit()
        return redirect(url_for("v2_complete.restocks"))
    query = V2RestockRequest.query
    if not is_admin(): query = query.filter_by(requested_by_id=current_user.id)
    rows = query.order_by(V2RestockRequest.requested_at.desc()).all()
    return render_template("v2/restocks.html", rows=rows)


@v2_complete_bp.route("/notifications")
@login_required
@roles_required("admin", "commercial", "animateur")
def smart_notifications():
    division = current_user.project if not is_admin() else _division_arg()
    divisions = [division] if division in DIVISION_SUPPLIERS else _divisions()
    alerts = []
    for row in V2Rupture.query.filter(V2Rupture.division.in_(divisions), V2Rupture.status != "resolved").order_by(V2Rupture.reported_at.desc()).limit(50):
        alerts.append(("Rupture", row.reported_at, f"{row.product_name} / {row.wholesaler or 'grossiste non précisé'}", "critical"))
    today = date.today()
    opp_q = V2Opportunity.query.filter(V2Opportunity.division.in_(divisions), V2Opportunity.next_followup <= today, V2Opportunity.next_followup.isnot(None))
    if not is_admin(): opp_q = opp_q.filter(V2Opportunity.owner_id == current_user.id)
    for row in opp_q.order_by(V2Opportunity.next_followup.asc()).limit(50):
        alerts.append(("Relance", datetime.combine(row.next_followup, datetime.min.time()), row.client.name, "warning"))
    exec_q = V2PlanningExecution.query.filter(
        V2PlanningExecution.execution_date < today, V2PlanningExecution.status == "planned"
    )
    if not is_admin(): exec_q = exec_q.filter_by(user_id=current_user.id)
    for row in exec_q.order_by(V2PlanningExecution.execution_date.desc()).limit(50):
        alerts.append(("Planning", datetime.combine(row.execution_date, datetime.min.time()), f"{row.structure_name} non clôturé", "warning"))
    return render_template("v2/notifications.html", alerts=sorted(alerts, key=lambda x: x[1], reverse=True))


@v2_complete_bp.route("/base/duplicates")
@login_required
@roles_required("admin")
def duplicates():
    groups = db.session.query(Client.name, Client.phone, func.count(Client.id)).group_by(
        Client.name, Client.phone
    ).having(func.count(Client.id) > 1).order_by(func.count(Client.id).desc()).all()
    rows = []
    for name, phone, count in groups:
        clients = Client.query.filter(Client.name == name, Client.phone == phone).order_by(Client.id).all()
        rows.append({"name": name, "phone": phone, "count": count, "clients": clients})
    return render_template("v2/duplicates.html", rows=rows)


@v2_complete_bp.route("/produits-performance")
@login_required
@roles_required("admin")
def product_performance():
    division = _division_arg()
    divisions = _divisions() if division == "all" else [division]
    rows = []
    for div in divisions:
        for slug in DIVISION_SUPPLIERS.get(div, []):
            supplier = SUPPLIERS[slug]
            product_model, sale_model = supplier["product_model"], supplier["sale_model"]
            products = product_model.query.filter_by(is_active=True).all()
            for product in products:
                quantity = db.session.query(func.coalesce(func.sum(sale_model.quantity), 0)).filter(
                    sale_model.product_id == product.id, sale_model.project == div
                ).scalar() or 0
                revenue = db.session.query(func.coalesce(func.sum(sale_model.quantity * sale_model.price), 0)).filter(
                    sale_model.product_id == product.id, sale_model.project == div
                ).scalar() or 0
                stocks = [getattr(product, f"stock_{w}", 0) or 0 for w in ("duopharm", "ubipharm", "laborex", "sodipharm")]
                rows.append({
                    "division": div, "laboratory": supplier["label"], "product": product.name,
                    "quantity": int(quantity), "revenue": Decimal(str(revenue or 0)),
                    "stock": sum(stocks), "rupture": sum(1 for x in stocks if x <= 0),
                })
    return render_template("v2/product_performance_complete.html", rows=rows)


@v2_complete_bp.route("/objectifs-multi", methods=["GET", "POST"])
@login_required
@roles_required("admin")
def multi_objectives():
    year = request.args.get("year", date.today().year, type=int)
    if request.method == "POST":
        year = request.form.get("year", year, type=int)
        division = (request.form.get("division") or "").strip().lower()
        metric = (request.form.get("metric") or "").strip().lower()
        month_raw = (request.form.get("month") or "").strip()
        user_id = request.form.get("user_id", type=int)
        target_raw = request.form.get("target_value", type=float)
        if division not in DIVISION_SUPPLIERS:
            flash("Division invalide.", "danger")
            return redirect(url_for("v2_complete.multi_objectives", year=year))
        if metric not in {"ca", "visites", "prospections", "prescriptions", "animations"}:
            flash("Indicateur invalide.", "danger")
            return redirect(url_for("v2_complete.multi_objectives", year=year))
        month = int(month_raw) if month_raw else None
        if month is not None and not 1 <= month <= 12:
            flash("Mois invalide.", "danger")
            return redirect(url_for("v2_complete.multi_objectives", year=year))
        if target_raw is None or target_raw < 0:
            flash("Objectif invalide.", "danger")
            return redirect(url_for("v2_complete.multi_objectives", year=year))
        user = User.query.get(user_id) if user_id else None
        if user and user.role not in ("commercial", "animateur"):
            flash("Utilisateur invalide.", "danger")
            return redirect(url_for("v2_complete.multi_objectives", year=year))
        if user and user.project != division:
            flash("La division de l'utilisateur ne correspond pas à l'objectif.", "danger")
            return redirect(url_for("v2_complete.multi_objectives", year=year))
        existing = V2Objective.query.filter_by(
            user_id=user.id if user else None,
            division=division, year=year, month=month, metric=metric
        ).first()
        if existing:
            existing.target_value = target_raw
        else:
            db.session.add(V2Objective(
                user_id=user.id if user else None,
                division=division, year=year, month=month, metric=metric,
                target_value=target_raw
            ))
        db.session.commit()
        flash("Objectif enregistré.", "success")
        return redirect(url_for("v2_complete.multi_objectives", year=year))

    rows = V2Objective.query.filter_by(year=year).order_by(
        V2Objective.division, V2Objective.month, V2Objective.metric
    ).all()
    users = User.query.filter(
        User.role.in_(("commercial", "animateur")),
        User.is_active_account.is_(True)
    ).order_by(User.project, User.username).all()
    return render_template("v2/objectives_multi.html", rows=rows, year=year, users=users)


def _report_rows(start, end, division):
    users = {u.id: u.username for u in User.query.all()}
    result = []
    for d in (_divisions() if division == "all" else [division]):
        for slug in DIVISION_SUPPLIERS.get(d, []):
            model = SUPPLIERS[slug]["sale_model"]
            amount = db.session.query(func.coalesce(func.sum(model.quantity * model.price), 0)).filter(
                model.date >= start, model.date < end, model.project == d
            ).scalar() or 0
            result.append([d, SUPPLIERS[slug]["label"], Decimal(str(amount))])
    return result


@v2_complete_bp.route("/rapports/automatique.xlsx")
@login_required
@roles_required("admin")
def automatic_xlsx():
    import xlsxwriter
    start = date.fromisoformat(request.args.get("start")) if request.args.get("start") else date.today().replace(day=1)
    end = date.fromisoformat(request.args.get("end")) if request.args.get("end") else date.today() + timedelta(days=1)
    division = _division_arg()
    output = io.BytesIO()
    workbook = xlsxwriter.Workbook(output, {"in_memory": True})
    sheet = workbook.add_worksheet("Synthèse")
    headers = ["Division", "Laboratoire", "CA"]
    for col, value in enumerate(headers): sheet.write(0, col, value)
    for row_idx, row in enumerate(_report_rows(start, end, division), 1):
        for col, value in enumerate(row): sheet.write(row_idx, col, value)
    sheet.write(row_idx + 2, 0, "Visites")
    sheet.write(row_idx + 2, 1, ClientVisit.query.filter(ClientVisit.date >= start, ClientVisit.date < end, ClientVisit.is_duplicate.is_(False)).count())
    workbook.close(); output.seek(0)
    return Response(output.read(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=nasora_v2_rapport.xlsx"})


def _report_summary(start, end, division):
    divisions = _divisions() if division == "all" else [division]
    return {
        "ca": sum((_revenue(d, start, end) for d in divisions), Decimal("0")),
        "visits": ClientVisit.query.filter(ClientVisit.date >= start, ClientVisit.date < end, ClientVisit.is_duplicate.is_(False)).count(),
        "prospections": Prospection.query.filter(Prospection.date >= start, Prospection.date < end).count(),
        "animations": AnimationSale.query.filter(AnimationSale.animation_date >= start, AnimationSale.animation_date < end).count(),
        "ruptures": V2Rupture.query.filter(V2Rupture.division.in_(divisions), V2Rupture.status != "resolved").count(),
    }


@v2_complete_bp.route("/rapports/automatique.pdf")
@login_required
@roles_required("admin")
def automatic_pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    start = date.fromisoformat(request.args.get("start")) if request.args.get("start") else date.today().replace(day=1)
    end = date.fromisoformat(request.args.get("end")) if request.args.get("end") else date.today() + timedelta(days=1)
    summary = _report_summary(start, end, _division_arg())
    output = io.BytesIO(); pdf = canvas.Canvas(output, pagesize=A4)
    y = 800; pdf.setFont("Helvetica-Bold", 16); pdf.drawString(50, y, "NASORA CRM V2 - Rapport automatique"); y -= 35
    pdf.setFont("Helvetica", 11)
    for label, value in (("Période", f"{start} au {end - timedelta(days=1)}"), ("CA", f"{summary['ca']} FCFA"), ("Visites", summary["visits"]), ("Prospections", summary["prospections"]), ("Animations", summary["animations"]), ("Ruptures ouvertes", summary["ruptures"])):
        pdf.drawString(50, y, f"{label}: {value}"); y -= 22
    pdf.save(); output.seek(0)
    return Response(output.read(), mimetype="application/pdf", headers={"Content-Disposition": "attachment; filename=nasora_v2_rapport.pdf"})


@v2_complete_bp.route("/rapports/automatique.docx")
@login_required
@roles_required("admin")
def automatic_docx():
    from docx import Document
    start = date.fromisoformat(request.args.get("start")) if request.args.get("start") else date.today().replace(day=1)
    end = date.fromisoformat(request.args.get("end")) if request.args.get("end") else date.today() + timedelta(days=1)
    summary = _report_summary(start, end, _division_arg())
    document = Document()
    document.add_heading("NASORA CRM V2 - Rapport automatique", 0)
    document.add_paragraph(f"Période : {start} au {end - timedelta(days=1)}")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Indicateur", "Valeur"
    for label, value in (("CA", f"{summary['ca']} FCFA"), ("Visites", summary["visits"]), ("Prospections", summary["prospections"]), ("Animations", summary["animations"]), ("Ruptures ouvertes", summary["ruptures"])):
        cells = table.add_row().cells; cells[0].text = label; cells[1].text = str(value)
    output = io.BytesIO(); document.save(output); output.seek(0)
    return Response(output.read(), mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": "attachment; filename=nasora_v2_rapport.docx"})


@v2_complete_bp.route("/offline")
@login_required
@roles_required("admin", "commercial", "animateur")
def offline():
    return render_template("v2/offline.html")
